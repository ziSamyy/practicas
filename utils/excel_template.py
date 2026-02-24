"""
Generación y lectura del Excel de configuración de la Orden de Transporte.

Endpoints que usan este módulo:
  GET  /ordenes/config/excel  → descarga el Excel con la config actual
  POST /ordenes/config/excel  → sube el Excel editado y actualiza el JSON
"""
import json
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Tuple

import openpyxl
from openpyxl.styles import (Alignment, Border, Font, PatternFill, Side)
from openpyxl.utils import get_column_letter

_CONFIG_PATH = Path(__file__).parent / "orden_template_config.json"

# ── colores de celda ──────────────────────────────────────────────────────────
_HEADER_FILL   = PatternFill("solid", fgColor="1F3864")   # azul oscuro
_SECTION_FILL  = PatternFill("solid", fgColor="2E75B6")   # azul medio
_LABEL_FILL    = PatternFill("solid", fgColor="D9E2F3")   # azul muy claro
_EDITABLE_FILL = PatternFill("solid", fgColor="FFF2CC")   # amarillo
_DESC_FILL     = PatternFill("solid", fgColor="F2F2F2")   # gris claro
_WHITE_FILL    = PatternFill("solid", fgColor="FFFFFF")

_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# ── definición de filas del template ─────────────────────────────────────────
# (seccion_json, clave_json, nombre_visible, descripcion, tipo)
# tipo: "text" | "number" | "percent" | "color"
ROWS: List[Tuple[str, str, str, str, str]] = [
    # ── Empresa ──────────────────────────────────────────────────────────────
    ("empresa", "nombre",
     "Nombre",
     "Nombre grande en el logo  (ej: ADK)",
     "text"),
    ("empresa", "subtitulo",
     "Subtítulo",
     "Texto pequeño bajo el nombre  (ej: MAQUINAS ENVASADORAS)",
     "text"),

    # ── Documento ─────────────────────────────────────────────────────────────
    ("documento", "titulo",
     "Título principal",
     "Título central del documento  (ej: ORDEN DE TRANSPORTE)",
     "text"),
    ("documento", "codigo",
     "Código / subtítulo",
     "Subtítulo central  (ej: ADK S.R.L.  RG 1.1)",
     "text"),
    ("documento", "revision",
     "Revisión",
     "Texto de revisión  (ej: Revisión: 00)",
     "text"),
    ("documento", "fecha_vigencia",
     "Fecha de vigencia",
     "Fecha fija del formato  (ej: Fecha de vigencia: 05/08/03)",
     "text"),
    ("documento", "hoja",
     "Hoja",
     "Número de hoja  (ej: Hoja: 1/1)",
     "text"),

    # ── Secciones ─────────────────────────────────────────────────────────────
    ("secciones", "datos_destinatario",
     "Barra 'Datos del destinatario'",
     "Encabezado de la sección de datos del cliente",
     "text"),
    ("secciones", "descripcion_orden",
     "Barra 'Descripción de la Orden'",
     "Encabezado de la sección de descripción",
     "text"),
    ("secciones", "valores_enviados",
     "Columna 'Valores enviados'",
     "Etiqueta de la primera columna de la tabla financiera",
     "text"),
    ("secciones", "banco",
     "Columna 'Banco'",
     "Encabezado de la columna Banco",
     "text"),
    ("secciones", "numero",
     "Columna 'Número'",
     "Encabezado de la columna Número de cheque/transferencia",
     "text"),
    ("secciones", "importe",
     "Columna 'Importe'",
     "Encabezado de la columna Importe",
     "text"),

    # ── Columna derecha ───────────────────────────────────────────────────────
    ("columna_derecha", "fila_1_label",
     "Etiqueta fila 1 (Valor Declarado)",
     "Ej: Valor Declarado",
     "text"),
    ("columna_derecha", "fila_2_label",
     "Etiqueta fila 2 (Efectivo enviado)",
     "Ej: Efectivo enviado:",
     "text"),
    ("columna_derecha", "fila_3_label",
     "Etiqueta fila 3 (Total)",
     "Ej: Total (Valores + Efvo.):",
     "text"),
    ("columna_derecha", "fila_4_label",
     "Etiqueta fila 4 (Flete)",
     "Ej: Flete a cargo de:",
     "text"),
    ("columna_derecha", "fila_5_label",
     "Etiqueta fila 5 (Transporte)",
     "Ej: Transporte:",
     "text"),

    # ── Pie de página ─────────────────────────────────────────────────────────
    ("pie_pagina", "texto_izquierda",
     "Pie — texto izquierda",
     "Ej: Emitio la Orden:",
     "text"),
    ("pie_pagina", "texto_centro",
     "Pie — texto centro",
     "Ej: Original",
     "text"),
    ("pie_pagina", "texto_derecha",
     "Pie — texto derecha",
     "Ej: PARA EL TRANSPORTE",
     "text"),

    # ── Colores ───────────────────────────────────────────────────────────────
    ("colores", "barra_seccion",
     "Color barra de sección",
     "RGB entre 0.0 y 1.0 separado por comas  (ej: 0.87, 0.87, 0.87 = gris claro  |  0,0.47,0.84 = azul)",
     "color"),

    # ── Dimensiones ───────────────────────────────────────────────────────────
    ("dimensiones", "margen_izquierdo",
     "Margen izquierdo (mm)",
     "Distancia desde el borde izquierdo de la hoja",
     "number"),
    ("dimensiones", "ancho_contenido",
     "Ancho del contenido (mm)",
     "Ancho total del formulario (A4 = 210 mm, recomendado 190)",
     "number"),
    ("dimensiones", "altura_encabezado",
     "Alto del encabezado (mm)",
     "Altura del bloque con logo + título",
     "number"),
    ("dimensiones", "alto_barra_seccion",
     "Alto barra de sección (mm)",
     "Altura de las barras grises de cada sección",
     "number"),
    ("dimensiones", "alto_fila_datos",
     "Alto fila de datos (mm)",
     "Altura de cada fila en 'Datos del destinatario'",
     "number"),
    ("dimensiones", "alto_descripcion",
     "Alto área de descripción (mm)",
     "Altura del área de texto libre  (más mm = más espacio para escribir)",
     "number"),
    ("dimensiones", "alto_fila_financiera",
     "Alto fila financiera (mm)",
     "Altura de cada fila en la sección de valores enviados",
     "number"),
    ("dimensiones", "filas_valores_enviados",
     "Cantidad de filas de valores enviados",
     "Número de filas de datos en la tabla (sin contar la cabecera)",
     "number"),
    ("dimensiones", "alto_pie",
     "Alto del pie (mm)",
     "Altura de la fila del pie de página",
     "number"),
    ("dimensiones", "ancho_caja_fecha",
     "Ancho caja fecha (mm)",
     "Ancho de la caja 'Fecha de envío' a la derecha",
     "number"),

    # ── Tabla financiera ──────────────────────────────────────────────────────
    ("tabla_financiera", "porcentaje_tabla_izquierda",
     "% ancho tabla izquierda",
     "Porcentaje del ancho total para la tabla  (0.60 = 60%,  resto va a columna derecha)",
     "percent"),
    ("tabla_financiera", "ancho_col_valores_enviados",
     "Ancho col. 'Valores enviados' (mm)",
     "Ancho de la primera columna (etiqueta)",
     "number"),
    ("tabla_financiera", "ancho_col_banco",
     "Ancho col. 'Banco' (mm)",
     "Ancho de la columna Banco",
     "number"),
    ("tabla_financiera", "ancho_col_numero",
     "Ancho col. 'Número' (mm)",
     "Ancho de la columna Número  (la col. Importe se calcula automático)",
     "number"),
    ("tabla_financiera", "porcentaje_label_derecha",
     "% etiqueta columna derecha",
     "Del ancho derecho, % destinado a etiquetas  (0.55 = 55%,  resto al valor)",
     "percent"),

    # ── Encabezado columnas ───────────────────────────────────────────────────
    ("encabezado_columnas", "porcentaje_logo",
     "% ancho celda logo",
     "Porcentaje del ancho total para la celda del logo  (0.263 = 26.3%)",
     "percent"),
    ("encabezado_columnas", "porcentaje_titulo",
     "% ancho celda título",
     "Porcentaje del ancho total para la celda del título  (resto va a info)",
     "percent"),

    # ── Pie columnas ──────────────────────────────────────────────────────────
    ("pie_columnas", "porcentaje_izquierda",
     "% pie — columna izquierda",
     "Porcentaje del pie para la sección 'Emitio la Orden'",
     "percent"),
    ("pie_columnas", "porcentaje_centro",
     "% pie — columna centro",
     "Porcentaje del pie para la sección central ('Original')",
     "percent"),
    ("pie_columnas", "porcentaje_derecha",
     "% pie — columna derecha",
     "Porcentaje del pie para 'PARA EL TRANSPORTE'",
     "percent"),

    # ── Fuentes ───────────────────────────────────────────────────────────────
    ("fuentes", "label_size",
     "Tamaño fuente — etiquetas",
     "Tamaño de texto para etiquetas pequeñas  (en puntos)",
     "number"),
    ("fuentes", "valor_size",
     "Tamaño fuente — valores",
     "Tamaño de texto para valores en negrita  (en puntos)",
     "number"),
    ("fuentes", "titulo_size",
     "Tamaño fuente — título principal",
     "Tamaño de texto del título central  (en puntos)",
     "number"),
    ("fuentes", "subtitulo_size",
     "Tamaño fuente — subtítulo",
     "Tamaño de texto del código/subtítulo  (en puntos)",
     "number"),
    ("fuentes", "descripcion_size",
     "Tamaño fuente — descripción",
     "Tamaño de texto en el área de descripción  (en puntos)",
     "number"),
    ("fuentes", "info_encabezado_size",
     "Tamaño fuente — info encabezado",
     "Tamaño de texto en la celda info (revisión / fecha vigencia / hoja)",
     "number"),
]

# Secciones visibles (para los encabezados de grupo)
_SECTION_LABELS: Dict[str, str] = {
    "empresa":              "🏢  EMPRESA",
    "documento":            "📄  DOCUMENTO",
    "secciones":            "🏷️  ETIQUETAS DE SECCIONES",
    "columna_derecha":      "📊  COLUMNA DERECHA (etiquetas)",
    "pie_pagina":           "📋  PIE DE PÁGINA",
    "colores":              "🎨  COLORES",
    "dimensiones":          "📐  DIMENSIONES (en mm)",
    "tabla_financiera":     "💰  TABLA FINANCIERA",
    "encabezado_columnas":  "📰  ENCABEZADO — PROPORCIONES",
    "pie_columnas":         "📏  PIE — PROPORCIONES",
    "fuentes":              "🔤  FUENTES (tamaño en puntos)",
}


# ─── helpers ──────────────────────────────────────────────────────────────────

def _value_to_str(section: str, key: str, cfg: Dict[str, Any]) -> str:
    """Convierte un valor del JSON a texto para la celda Excel."""
    val = cfg.get(section, {}).get(key)
    if isinstance(val, list):           # color RGB
        return ", ".join(str(v) for v in val)
    if val is None:
        return ""
    return str(val)


def _str_to_value(raw: str, tipo: str):
    """Convierte el texto de la celda Excel al tipo correcto para el JSON."""
    raw = str(raw).strip()
    if tipo == "number":
        try:
            f = float(raw)
            return int(f) if f == int(f) else f
        except ValueError:
            return raw
    if tipo == "percent":
        try:
            return float(raw)
        except ValueError:
            return raw
    if tipo == "color":
        parts = [p.strip() for p in raw.replace(";", ",").split(",")]
        try:
            return [float(p) for p in parts if p]
        except ValueError:
            return raw
    return raw   # text


def _cell_style(ws, cell, fill, font_color="000000", bold=False,
                italic=False, align="left", wrap=False):
    cell.fill = fill
    cell.font = Font(bold=bold, italic=italic, color=font_color,
                     name="Calibri", size=10)
    cell.alignment = Alignment(horizontal=align, vertical="center",
                                wrap_text=wrap)
    cell.border = _BORDER


# ─── generación del Excel ─────────────────────────────────────────────────────

def generar_excel_template() -> bytes:
    """Lee el JSON de configuración y genera un Excel editable."""
    with open(_CONFIG_PATH, encoding="utf-8") as f:
        cfg = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Configuración"

    # ── fila de título ────────────────────────────────────────────────────────
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "⚙️  CONFIGURACIÓN DE LA ORDEN DE TRANSPORTE  —  Editá la columna 'VALOR' y subí el archivo"
    _cell_style(ws, title_cell, _HEADER_FILL, font_color="FFFFFF",
                bold=True, align="center")
    title_cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=12)
    ws.row_dimensions[1].height = 28

    # ── cabecera de columnas ──────────────────────────────────────────────────
    headers = ["SECCIÓN", "CAMPO", "VALOR  ✏️  (editá aquí)", "DESCRIPCIÓN"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        _cell_style(ws, cell, _SECTION_FILL, font_color="FFFFFF",
                    bold=True, align="center")
    ws.row_dimensions[2].height = 20

    # Columnas ocultas con claves internas (para el parser)
    ws.cell(row=2, column=5, value="__section__")
    ws.cell(row=2, column=6, value="__key__")
    ws.cell(row=2, column=7, value="__type__")
    ws.column_dimensions["E"].hidden = True
    ws.column_dimensions["F"].hidden = True
    ws.column_dimensions["G"].hidden = True

    # ── filas de datos ────────────────────────────────────────────────────────
    current_section = None
    excel_row = 3

    for section, key, label, desc, tipo in ROWS:

        # separador de sección
        if section != current_section:
            current_section = section
            ws.merge_cells(f"A{excel_row}:D{excel_row}")
            sc = ws.cell(row=excel_row, column=1,
                         value=_SECTION_LABELS.get(section, section.upper()))
            _cell_style(ws, sc, _SECTION_FILL, font_color="FFFFFF",
                        bold=True, align="left")
            ws.row_dimensions[excel_row].height = 18
            excel_row += 1

        value_str = _value_to_str(section, key, cfg)

        # columna A — sección (vacía en filas de datos)
        ca = ws.cell(row=excel_row, column=1, value="")
        _cell_style(ws, ca, _WHITE_FILL)

        # columna B — nombre del campo
        cb = ws.cell(row=excel_row, column=2, value=label)
        _cell_style(ws, cb, _LABEL_FILL, bold=False)

        # columna C — valor EDITABLE
        cc = ws.cell(row=excel_row, column=3, value=value_str)
        _cell_style(ws, cc, _EDITABLE_FILL, bold=True, align="left")

        # columna D — descripción
        cd = ws.cell(row=excel_row, column=4, value=desc)
        _cell_style(ws, cd, _DESC_FILL, italic=True, wrap=True)

        # columnas ocultas (para el parser)
        ws.cell(row=excel_row, column=5, value=section)
        ws.cell(row=excel_row, column=6, value=key)
        ws.cell(row=excel_row, column=7, value=tipo)

        ws.row_dimensions[excel_row].height = 18
        excel_row += 1

    # ── anchos de columna ─────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 62

    # Fijar las primeras 2 filas
    ws.freeze_panes = "A3"

    # ── instrucciones en una segunda hoja ─────────────────────────────────────
    ws2 = wb.create_sheet("📖 Instrucciones")
    instrucciones = [
        ("¿CÓMO USAR ESTE ARCHIVO?", True, "1F3864", "FFFFFF"),
        ("", False, "FFFFFF", "000000"),
        ("1. Editá SOLO la columna 'VALOR' en la hoja 'Configuración'.", False, "FFFFFF", "000000"),
        ("2. NO cambies las columnas Sección, Campo ni Descripción.", False, "FFF2CC", "000000"),
        ("3. NO agregues ni elimines filas.", False, "FFF2CC", "000000"),
        ("4. Guardá el archivo como .xlsx (no cambies el formato).", False, "FFFFFF", "000000"),
        ("5. Subí el archivo al endpoint:  POST /ordenes/config/excel", False, "FFFFFF", "000000"),
        ("", False, "FFFFFF", "000000"),
        ("TIPOS DE VALORES", True, "2E75B6", "FFFFFF"),
        ("", False, "FFFFFF", "000000"),
        ("  texto     →  Escribí cualquier texto. Ej: ORDEN DE TRANSPORTE", False, "FFFFFF", "000000"),
        ("  número    →  Escribí un número entero o decimal. Ej: 10  |  8.5", False, "FFFFFF", "000000"),
        ("  porcentaje→  Escribí un decimal entre 0 y 1. Ej: 0.60 (= 60%)", False, "FFFFFF", "000000"),
        ("  color     →  Tres valores RGB de 0.0 a 1.0 separados por coma.", False, "FFFFFF", "000000"),
        ("             Ej: 0.87, 0.87, 0.87  (gris claro)", False, "D9E2F3", "000000"),
        ("             Ej: 0, 0.47, 0.84     (azul)", False, "D9E2F3", "000000"),
        ("             Ej: 1, 1, 1            (blanco)", False, "D9E2F3", "000000"),
        ("             Ej: 0, 0, 0            (negro)", False, "D9E2F3", "000000"),
    ]
    for i, (texto, bold, bg, fg) in enumerate(instrucciones, 1):
        ws2.merge_cells(f"A{i}:D{i}")
        cell = ws2.cell(row=i, column=1, value=texto)
        cell.fill = PatternFill("solid", fgColor=bg)
        cell.font = Font(bold=bold, color=fg, name="Calibri", size=11)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws2.row_dimensions[i].height = 20
    ws2.column_dimensions["A"].width = 80

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ─── lectura / parse del Excel subido ─────────────────────────────────────────

def leer_excel_y_actualizar_config(file_bytes: bytes) -> Dict[str, Any]:
    """
    Lee el Excel subido, extrae los valores editados y
    actualiza el JSON de configuración.
    Devuelve el config actualizado como dict.
    """
    wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb["Configuración"]

    # Cargar config actual (para no pisar claves no mapeadas)
    with open(_CONFIG_PATH, encoding="utf-8") as f:
        cfg: Dict[str, Any] = json.load(f)

    for row in ws.iter_rows(min_row=3, values_only=True):
        # columnas ocultas E, F, G (índices 4, 5, 6)
        if len(row) < 7:
            continue
        section = row[4]
        key     = row[5]
        tipo    = row[6]
        value   = row[2]   # columna C (índice 2)

        if not section or not key or section.startswith("__"):
            continue

        # Saltear filas de encabezado de sección (sin key útil)
        if str(section).startswith("🏢") or str(key) == "__key__":
            continue

        if section not in cfg:
            cfg[section] = {}

        parsed = _str_to_value(str(value) if value is not None else "", tipo)
        cfg[section][key] = parsed

    # Guardar el JSON actualizado
    with open(_CONFIG_PATH, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)

    return cfg
